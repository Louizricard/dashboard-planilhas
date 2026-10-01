import csv
import hashlib
import io
import json
import os
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock

import openpyxl
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from .domain import FIELDS, validate, external

ROOT = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get('FRUTMAR_DB', ROOT.parent / '.local/frutmar.sqlite3'))
app = FastAPI(title='Frutmar Data Hub', version='0.1.0')
app.mount('/static', StaticFiles(directory=ROOT / 'static'), name='static')
previews = {}
preview_lock = Lock()

def now(): return datetime.now(timezone.utc).isoformat()

@contextmanager
def database():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    conn.executescript('''
    CREATE TABLE IF NOT EXISTS products (
        id TEXT PRIMARY KEY, code TEXT UNIQUE, name TEXT NOT NULL,
        category TEXT NOT NULL, payload TEXT NOT NULL,
        source TEXT NOT NULL, updated_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS imports (
        id TEXT PRIMARY KEY, filename TEXT NOT NULL, hash TEXT UNIQUE NOT NULL,
        created_at TEXT NOT NULL, inserted INTEGER NOT NULL, errors TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS audit (
        id INTEGER PRIMARY KEY, product_id TEXT NOT NULL, action TEXT NOT NULL,
        before_value TEXT, after_value TEXT NOT NULL, source TEXT NOT NULL, created_at TEXT NOT NULL);
    ''')
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally: conn.close()

def insert(conn, p, source):
    payload = json.dumps(p, ensure_ascii=False)
    conn.execute('INSERT INTO products VALUES (?,?,?,?,?,?,?)',
                 (p['id'], p['codigoSistema'], p['nome'], p['categoria'], payload, source, now()))
    conn.execute('INSERT INTO audit(product_id,action,after_value,source,created_at) VALUES (?,?,?,?,?)',
                 (p['id'], 'create', payload, source, now()))

@app.get('/')
def home(): return FileResponse(ROOT / 'static/index.html')

@app.get('/api/products')
def products():
    with database() as conn:
        return [external(json.loads(row['payload'])) | {'origem':row['source'], 'atualizadoEm':row['updated_at']}
                for row in conn.execute('SELECT * FROM products ORDER BY name')]

@app.post('/api/products', status_code=201)
def create_product(raw: dict):
    try: p = validate(raw)
    except ValueError as e: raise HTTPException(422, str(e))
    try:
        with database() as conn: insert(conn, p, 'cadastro manual')
    except sqlite3.IntegrityError: raise HTTPException(409, 'ID ou código do sistema já cadastrado')
    return external(p)

@app.put('/api/products/{product_id}')
def edit_product(product_id: str, raw: dict):
    if raw.get('id') != product_id: raise HTTPException(422, 'O identificador do produto é imutável')
    try: p = validate(raw)
    except ValueError as e: raise HTTPException(422, str(e))
    try:
        with database() as conn:
            old = conn.execute('SELECT payload FROM products WHERE id=?', (product_id,)).fetchone()
            if not old: raise HTTPException(404, 'Produto não encontrado')
            payload = json.dumps(p, ensure_ascii=False)
            conn.execute('UPDATE products SET code=?,name=?,category=?,payload=?,updated_at=? WHERE id=?',
                         (p['codigoSistema'],p['nome'],p['categoria'],payload,now(),product_id))
            conn.execute('INSERT INTO audit(product_id,action,before_value,after_value,source,created_at) VALUES (?,?,?,?,?,?)',
                         (product_id,'update',old['payload'],payload,'edição manual',now()))
    except sqlite3.IntegrityError: raise HTTPException(409, 'Código do sistema já cadastrado')
    return external(p)

@app.get('/api/products/{product_id}/history')
def history(product_id: str):
    with database() as conn:
        return [dict(r) for r in conn.execute('SELECT * FROM audit WHERE product_id=? ORDER BY id DESC',(product_id,))]

@app.post('/api/imports/preview')
async def preview(file: UploadFile):
    data = await file.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024: raise HTTPException(413, 'Limite de arquivo: 5 MB')
    digest = hashlib.sha256(data).hexdigest()
    with database() as conn:
        if conn.execute('SELECT 1 FROM imports WHERE hash=?',(digest,)).fetchone():
            raise HTTPException(409,'Este arquivo já foi importado')
        existing_ids = {r['id'] for r in conn.execute('SELECT id FROM products')}
        existing_codes = {r['code'] for r in conn.execute('SELECT code FROM products WHERE code IS NOT NULL')}
    try:
        suffix = Path(file.filename or '').suffix.lower()
        if suffix == '.xlsx':
            import zipfile
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Arquivo descompactado excede 50 MB')
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            if 'Produtos' not in book.sheetnames: raise ValueError('Aba Produtos não encontrada')
            sheet = book['Produtos']
            if sheet.max_row > 10001: raise ValueError('Limite de 10.000 linhas')
            values = list(sheet.values)
            book.close()
            headers, rows = values[0], values[1:]
        elif suffix == '.csv':
            content = data.decode('utf-8-sig')
            try: dialect = csv.Sniffer().sniff(content[:4096], delimiters=',;')
            except csv.Error: dialect = csv.excel
            values = list(csv.reader(io.StringIO(content), dialect))
            headers, rows = values[0], values[1:]
        else: raise ValueError('Use .xlsx ou .csv (UTF-8)')
        headers = [str(h or '').strip() for h in headers]
        if len(set(headers)) != len(headers): raise ValueError('Há cabeçalhos duplicados')
        missing = {'id','nome','categoria','unidadeVenda'} - set(headers)
        if missing: raise ValueError('Colunas ausentes: ' + ', '.join(sorted(missing)))
        if len(rows) > 10000: raise ValueError('Limite de 10.000 linhas')
    except Exception as e: raise HTTPException(422, f'Não foi possível ler o arquivo: {e}')
    valid, errors = [], []
    for index, row in enumerate(rows, 2):
        if not any(v is not None and v != '' for v in row): continue
        try:
            p = validate(dict(zip(headers,row)))
            if p['id'] in existing_ids or (p['codigoSistema'] and p['codigoSistema'] in existing_codes):
                raise ValueError('ID ou código duplicado; use a edição para atualizar um produto')
            existing_ids.add(p['id'])
            if p['codigoSistema']: existing_codes.add(p['codigoSistema'])
            valid.append(p)
        except ValueError as e: errors.append({'linha':index, 'erro':str(e)})
    token = str(uuid.uuid4())
    with preview_lock:
        # Previews are temporary, bounded and expire after 30 minutes.
        import time
        for key in list(previews):
            if time.time() - previews[key]['time'] > 1800: del previews[key]
        if len(previews) >= 20: raise HTTPException(429,'Muitas prévias abertas; tente novamente mais tarde')
        previews[token] = {'data':valid,'errors':errors,'hash':digest,'filename':Path(file.filename or 'arquivo').name,'time':time.time()}
    return {'token':token,'validos':len(valid),'erros':errors,'amostra':[external(p) for p in valid[:10]],'total':len(valid)+len(errors)}

@app.post('/api/imports/{token}/confirm')
def confirm(token: str, options: dict):
    import time
    with preview_lock:
        item = previews.get(token)
        if not item or time.time() - item['time'] > 1800: raise HTTPException(404,'Prévia expirada; envie o arquivo novamente')
        if item['errors'] and not options.get('permitirParcial', False):
            raise HTTPException(422,'Confirme explicitamente a importação parcial')
        if not item['data']: raise HTTPException(422,'Nenhuma linha válida para importar')
        identifier = str(uuid.uuid4())
        try:
            with database() as conn:
                conn.execute('INSERT INTO imports VALUES (?,?,?,?,?,?)',
                             (identifier,item['filename'],item['hash'],now(),len(item['data']),json.dumps(item['errors'],ensure_ascii=False)))
                for p in item['data']: insert(conn,p,'importação '+identifier)
        except sqlite3.IntegrityError:
            raise HTTPException(409,'Os dados mudaram desde a prévia ou o arquivo já foi importado. Gere uma nova prévia.')
        del previews[token]
    return {'id':identifier,'importados':len(item['data'])}

@app.get('/api/imports')
def imports():
    with database() as conn:
        return [dict(r) | {'errors':json.loads(r['errors'])} for r in conn.execute('SELECT * FROM imports ORDER BY created_at DESC')]

@app.get('/api/export')
def export():
    output = io.StringIO()
    writer = csv.DictWriter(output,fieldnames=FIELDS,delimiter=';')
    writer.writeheader()
    for product in products():
        row = {k:product.get(k) for k in FIELDS}
        for k,v in row.items():
            if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')): row[k] = "'" + v
        writer.writerow(row)
    return Response('\ufeff'+output.getvalue(),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename=produtos-frutmar.csv'})
